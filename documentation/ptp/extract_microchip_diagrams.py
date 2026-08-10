#!/usr/bin/env python3
"""Extract relevant ACMA / PTP / TSU diagrams from the Microchip PPTX
templates and save them as PNG/JPEG files for reuse.

Output dir: documentation/pdf/_extracted_diagrams/
"""
import os
from pptx import Presentation
from pptx.util import Emu

REPO = r"c:\work\ptp\check4\net_10base_t1s"
OUT_DIR = os.path.join(REPO, "documentation", "pdf", "_extracted_diagrams")
os.makedirs(OUT_DIR, exist_ok=True)

# (source_pptx, slide_index_1based, output_filename, description)
TARGETS = [
    # From "3_Special Features of LAN8670_1_2 LAN8650_1.pptx"
    ("3_Special Features of LAN8670_1_2 LAN8650_1.pptx", 28,
     "ms_28_lan8650_acma_block.png",
     "LAN8650 ACMA block diagram with internal/external source"),
    ("3_Special Features of LAN8670_1_2 LAN8650_1.pptx", 33,
     "ms_33_better_sw_architecture.png",
     "Software architecture for periodic wall-clock updates"),
    ("3_Special Features of LAN8670_1_2 LAN8650_1.pptx", 37,
     "ms_37_rx_timestamp_config.png",
     "LAN8650/1 Receive Timestamp Configuration (pattern matcher)"),
    ("3_Special Features of LAN8670_1_2 LAN8650_1.pptx", 43,
     "ms_43_lan867x_8021as_config.png",
     "IEEE 802.1AS LAN8670/1/2 Standard Clock Synchronization config"),
    # From LAN867x_PHY_d47.pptx
    ("LAN867x_PHY_d47.pptx", 27,
     "phy_27_acma_config_example.png",
     "Configuration example for ACMA (LAN8670/1/2 family)"),
    ("LAN867x_PHY_d47.pptx", 34,
     "phy_34_usecases_standards.png",
     "Time-sync use cases and IEEE standards overview"),
    ("LAN867x_PHY_d47.pptx", 35,
     "phy_35_8021as_support.png",
     "IEEE 802.1AS Support — LAN8670/1/2 and LAN8650/1 family"),
]


def extract_pictures_from_slide(slide):
    """Return list of (image_blob, ext, dimensions_emu) for all
    pictures on the slide (including those inside groups)."""
    out = []

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == 13:  # PICTURE
                img = shape.image
                out.append((img.blob, img.ext, (shape.left, shape.top,
                                                 shape.width, shape.height)))
            elif shape.shape_type == 6:  # GROUP
                walk(shape.shapes)
    walk(slide.shapes)
    return out


def main():
    sources = {}
    for source, slide_idx, out_name, desc in TARGETS:
        if source not in sources:
            sources[source] = Presentation(
                os.path.join(REPO, "documentation", "pdf", source))
        prs = sources[source]
        if slide_idx > len(prs.slides):
            print(f"  SKIP: slide {slide_idx} > {len(prs.slides)} in {source}")
            continue
        slide = prs.slides[slide_idx - 1]
        pics = extract_pictures_from_slide(slide)
        if not pics:
            print(f"  SKIP: no pictures on {source} slide {slide_idx}")
            continue
        # Save the largest picture as the primary diagram
        largest = max(pics, key=lambda p: p[2][2] * p[2][3])
        blob, ext, dims = largest
        # Adjust filename extension to match actual format
        base, _ = os.path.splitext(out_name)
        out_path = os.path.join(OUT_DIR, base + "." + ext)
        with open(out_path, "wb") as f:
            f.write(blob)
        print(f"  EXTRACTED: {out_path}  ({len(pics)} pic(s), "
              f"saved largest = {dims[2]}x{dims[3]} EMU)  -- {desc}")


if __name__ == "__main__":
    main()
